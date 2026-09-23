#!/usr/bin/env python3
"""Aniverse V1→V2 발표 PPT — 사용자 확정본(10장) 기준.

기준 문서: CONTENT_V1V2_working.md
소스 PPT: ppt/Aniverse_V1V2_발표_working.pptx (업로드 확정본 구조)

출력: ppt/Aniverse_V1V2_발표_vN.pptx 및 working.pptx 갱신
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "맑은 고딕"
BASE = Path(__file__).resolve().parent
IMG = BASE / "images" / "hybrid"
OUT = BASE / "ppt"
OUT.mkdir(parents=True, exist_ok=True)

SW, SH = 13.333, 7.5
STEM = "Aniverse_V1V2_발표"
TOTAL = 10
WORKING = OUT / "Aniverse_V1V2_발표_working.pptx"

NAVY = RGBColor(15, 23, 42)
BLACK = RGBColor(17, 24, 39)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(107, 114, 128)
LIGHT = RGBColor(243, 244, 246)
SOFT = RGBColor(249, 250, 251)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(14, 165, 233)
RED = RGBColor(239, 68, 68)
GREEN = RGBColor(34, 197, 94)
ORANGE = RGBColor(249, 115, 22)
PURPLE = RGBColor(139, 92, 246)
TEAL = RGBColor(20, 184, 166)


def next_path() -> Path:
    best = 0
    for p in OUT.glob(f"{STEM}_v*.pptx"):
        m = re.search(r"_v(\d+)\.pptx$", p.name)
        if m:
            best = max(best, int(m.group(1)))
    return OUT / f"{STEM}_v{best + 1}.pptx"


def font(run, size=18, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = rPr.makeelement(qn(f"a:{tag}"), {})
            rPr.append(el)
        el.set("typeface", FONT)


def pad(tf, l=0.16, t=0.12, r=0.16, b=0.12):
    tf.word_wrap = True
    tf.margin_left = Inches(l)
    tf.margin_right = Inches(r)
    tf.margin_top = Inches(t)
    tf.margin_bottom = Inches(b)


def set_text(tf, text, size=18, bold=False, color=BLACK, align=None):
    tf.clear()
    pad(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    if align is not None:
        p.alignment = align


def add_para(tf, text, size=16, bold=False, color=BLACK, space_before=8, align=None):
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    p.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    return p


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def card(slide, x, y, w, h, fill=SOFT, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    else:
        sh.line.fill.background()
    pad(sh.text_frame)
    return sh


def circle_icon(slide, x, y, size, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def title_center(slide, text, y=0.22, size=36):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(y), Inches(12.3), Inches(0.7))
    set_text(box.text_frame, text, size, True, BLACK, PP_ALIGN.CENTER)


def footer(slide, n):
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(12.3), Inches(0.28))
    p = tx.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Aniverse  ·  Architecture V1 → V2"
    font(r, 13, False, GRAY)
    num = slide.shapes.add_textbox(Inches(11.3), Inches(7.15), Inches(1.5), Inches(0.28))
    p = num.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"{n} / {TOTAL}"
    font(r, 13, True, BLUE)
    p.alignment = PP_ALIGN.RIGHT


def put_img(slide, name, top=Inches(1.05), bottom=Inches(7.0), side=0.45):
    path = IMG / name
    if not path.exists():
        return False
    iw, ih = Image.open(path).size
    aspect = ih / float(iw)
    max_w = SW - side * 2
    max_h = bottom.inches - top.inches
    w_in = max_w
    h_in = w_in * aspect
    if h_in > max_h:
        h_in = max_h
        w_in = h_in / aspect
    x_in = (SW - w_in) / 2.0
    y_in = top.inches + (max_h - h_in) / 2.0
    slide.shapes.add_picture(
        str(path), Inches(x_in), Inches(y_in), width=Inches(w_in), height=Inches(h_in)
    )
    return True


def accent_item(slide, x, y, w, h, color, title, body):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.1), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    box = slide.shapes.add_textbox(x + Inches(0.22), y, w - Inches(0.3), h)
    tf = box.text_frame
    set_text(tf, title, 22, True, BLACK)
    add_para(tf, body, 18, False, GRAY, 10)


# ---------------------------------------------------------------------------
def slide_cover(prs):
    s = blank(prs)
    bg(s, RGBColor(239, 246, 255))
    t = s.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(8), Inches(2.5))
    tf = t.text_frame
    set_text(tf, "Aniverse", 54, True, BLUE)
    add_para(tf, "통합 서브컬처 커뮤니티 사이트", 24, False, NAVY, 14)
    add_para(tf, "Architecture V1 → V2  ·  EKS · GitOps 전환", 20, False, GRAY, 12)
    add_para(tf, "https://aniverse.my", 18, False, CYAN, 18)

    box = s.shapes.add_textbox(Inches(7.5), Inches(5.2), Inches(5.2), Inches(1.5))
    tf = box.text_frame
    set_text(tf, "Team", 18, True, BLUE, PP_ALIGN.RIGHT)
    add_para(tf, "김현우, 박서이, 김윤주, 강유민", 16, False, NAVY, 8, PP_ALIGN.RIGHT)


def slide_v1_def(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "Architecture V1 정의")
    items = [
        (
            "V1 정의",
            "User → ALB → ASG/EC2(Nginx+Django) → RDS · EFS · S3 로 Aniverse 서비스를 운영하던 초기 클라우드 구성입니다.",
        ),
        (
            "구현 방식",
            "Terraform으로 VPC·ALB·ASG를 구성하고, CodeDeploy / Actions로 EC2에 배포하는 방식을 사용했습니다.",
        ),
        (
            "개발 배경",
            "통합 서브컬처 사이트를 안정적으로 서비스하기 위해, 기존 온프레미스 구조에서 AWS 로 이전, EC2 기반 3-tier로 구축했습니다.",
        ),
        (
            "구현 범위",
            "HTTPS(ACM)·WAF·Redis·미디어 S3까지 포함하여 운영 가능한 형태까지 구현이 완료된 상태였습니다.",
        ),
    ]
    positions = [(0.55, 1.25), (6.85, 1.25), (0.55, 4.15), (6.85, 4.15)]
    for (x, y), (title, body) in zip(positions, items):
        c = card(s, Inches(x), Inches(y), Inches(5.95), Inches(2.6), LIGHT)
        set_text(c.text_frame, title, 22, True, BLUE)
        add_para(c.text_frame, body, 18, False, NAVY, 12)
    footer(s, 2)


def slide_arch_v1(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "시스템 구성도  ·  Architecture V1", size=34)
    put_img(s, "hybrid_01_arch_v1_vpc.png")
    footer(s, 3)


def slide_v1_to_v2(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "V1의 한계  ·  V2 목표  ·  V2 차별점", size=34)
    headers = [
        (0.4, RED, "V1의 한계"),
        (4.65, GREEN, "V2의 기능 · 개발 목표"),
        (8.9, BLUE, "V2 차별점"),
    ]
    for x, color, text in headers:
        h = card(s, Inches(x), Inches(1.05), Inches(4.05), Inches(0.65), color)
        set_text(h.text_frame, text, 20, True, WHITE, PP_ALIGN.CENTER)

    left = [
        ("모니터링이 한눈에 안 보임", "메트릭·로그·배포 상태가 흩어져 장애/부하를 한 화면에서 보기 어려웠습니다."),
        ("수정 → 서비스 반영이 어려움", "인프라·앱 변경이 EC2/에이전트에 묶여 서비스에 빠르고 일관되게 반영되기 어려웠습니다."),
        ("ASG만으로 늘렸다 줄이기", "스케일이 인스턴스 단위라, 필요할 때만 줄이기 어렵고 상시 용량으로 비용이 낭비될 수 있었습니다."),
    ]
    mid = [
        ("관측을 한곳에", "Prom/Grafana·Loki로 메트릭·로그를 모아 클러스터·앱 상태를 한눈에 봅니다."),
        ("변경하면 즉시 배포", "이미지 빌드 + GitOps(Argo)로 수정이 Synced 상태로 서비스에 반영됩니다."),
        ("필요한 만큼만 스케일", "노드·파드 단위로 스케일을 나누고 start/stop으로 미사용 비용을 줄입니다."),
    ]
    right = [
        ("관측 스택", "Prometheus · Grafana · Loki를 기본 구성에 두고, Alert/Tempo로 확장합니다."),
        ("Actions → ECR → Argo", "sha-* 태그 bump 후 Argo sync로 코드/차트 변경이 선언적으로 반영됩니다."),
        ("EKS + DB Pod", "ASG 인스턴스 스케일 대신 EKS 노드/파드 스케일과 PVC 기반 DB로 낭비를 줄입니다."),
    ]
    for i, (t, b) in enumerate(left):
        accent_item(s, Inches(0.4), Inches(1.85) + i * Inches(1.65), Inches(4.05), Inches(1.5), RED, t, b)
    for i, (t, b) in enumerate(mid):
        accent_item(s, Inches(4.65), Inches(1.85) + i * Inches(1.65), Inches(4.05), Inches(1.5), GREEN, t, b)
    for i, (t, b) in enumerate(right):
        accent_item(s, Inches(8.9), Inches(1.85) + i * Inches(1.65), Inches(4.05), Inches(1.5), BLUE, t, b)
    footer(s, 4)


def slide_v2_def(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "Architecture V2 정의")
    vision = card(s, Inches(0.55), Inches(1.15), Inches(12.2), Inches(1.45), LIGHT, BLUE)
    set_text(vision.text_frame, "비전", 20, True, BLUE)
    add_para(
        vision.text_frame,
        "Aniverse를 EKS 위에서 GitOps로 재현 가능하게 운영하고, 데이터·관측·보안까지 한 사이클로 다루는 최종 아키텍처를 목표로 합니다.",
        18,
        False,
        NAVY,
        8,
    )
    label = s.shapes.add_textbox(Inches(0.55), Inches(2.8), Inches(4), Inches(0.4))
    set_text(label.text_frame, "주요 기능", 20, True, BLACK)
    feats = [
        (BLUE, "ALB Ingress · HTTPS", "ACM ISSUED 후 Ingress ALB + Route53으로 외부 진입을 구성합니다."),
        (GREEN, "EKS web / db Pod", "Django+Nginx Deployment와 MariaDB StatefulSet+PVC로 앱·DB를 운영합니다."),
        (ORANGE, "Actions → ECR → Argo", "OIDC 빌드, sha-* 이미지, values bump 후 Argo sync로 배포합니다."),
        (PURPLE, "백업 · 관측", "restore Job / CronJob→S3, Prom·Grafana·Loki로 운영 가시성을 확보합니다."),
    ]
    for i, (color, title, body) in enumerate(feats):
        col, row = i % 2, i // 2
        x = Inches(0.55) + col * Inches(6.25)
        y = Inches(3.3) + row * Inches(1.7)
        card(s, x, y, Inches(6.0), Inches(1.5), SOFT)
        circle_icon(s, x + Inches(0.2), y + Inches(0.4), Inches(0.55), color)
        box = s.shapes.add_textbox(x + Inches(0.95), y + Inches(0.2), Inches(4.8), Inches(1.2))
        set_text(box.text_frame, title, 20, True, BLACK)
        add_para(box.text_frame, body, 17, False, GRAY, 8)
    footer(s, 5)


def slide_arch_v2(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "시스템 구성도  ·  Architecture V2", size=34)
    put_img(s, "hybrid_02_arch_v2_vpc.png")
    footer(s, 6)


def slide_v2_strengths(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "Architecture V2의 강점")
    rows = [
        ("배포 재현", "동일 이미지·매니페스트로 환경을 다시 올릴 것   →   Actions → ECR → Argo"),
        ("스케일 분리", "노드 스케일과 파드 HPA를 분리 설계   →   EKS + HPA 여지"),
        ("DB 비용", "RDS 상시 대신 학습용 영구 볼륨   →   StatefulSet + PVC"),
        ("보안", "장기 키 금지 · 이관 시 ARN 전수 교체   →   GitHub OIDC · IRSA"),
        ("관측", "메트릭·로그를 기본 구성에 포함   →   Prom / Grafana / Loki"),
    ]
    for i, (a, b) in enumerate(rows):
        y = Inches(1.2) + i * Inches(1.05)
        c = card(s, Inches(0.6), y, Inches(12.1), Inches(0.92), LIGHT if i % 2 == 0 else SOFT)
        set_text(c.text_frame, a, 18, True, BLUE)
        add_para(c.text_frame, b, 16, False, NAVY, 4)
    footer(s, 7)


def slide_verify(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "기능 · 운영 검증 기준")
    items = [
        ("웹 health", "Ingress/ALB 통해 HTTPS 200", "curl / 브라우저", "충족"),
        ("DB 시드", "restore Job 후 목록 UI에 데이터", "minTables + seed_rows", "충족"),
        ("GitOps", "image.tag bump → Argo Synced", "Actions 로그 + Argo UI", "충족"),
        ("미디어", "destroy 후 Sync media 재업로드", "S3 media/ 확인", "충족"),
        ("관측", "Prom/Grafana 대시보드 조회", "클러스터 메트릭", "구성"),
        ("알림/트레이싱", "AlertManager · Tempo", "다음 단계", "예정"),
    ]
    for i, (t, check, method, result) in enumerate(items):
        col, row = i % 3, i // 3
        x = Inches(0.45) + col * Inches(4.25)
        y = Inches(1.25) + row * Inches(2.7)
        c = card(s, x, y, Inches(4.05), Inches(2.45), LIGHT)
        set_text(c.text_frame, t, 20, True, BLUE)
        add_para(c.text_frame, "검증: " + check, 17, False, NAVY, 10)
        add_para(c.text_frame, "방법: " + method, 16, False, GRAY, 8)
        add_para(c.text_frame, "결과: " + result, 17, True, GREEN if result == "충족" else ORANGE, 10)
    footer(s, 8)


def slide_stack(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "기술 스택  ·  데이터 흐름")
    stack = card(s, Inches(0.5), Inches(1.15), Inches(5.9), Inches(5.5), LIGHT)
    set_text(stack.text_frame, "Tech Stack", 22, True, BLUE)
    for line in [
        "App: Django + Nginx",
        "Container: Docker · Helm",
        "Orchestration: EKS (V2)",
        "DB: MariaDB StatefulSet + EBS PVC",
        "Storage: S3 (media / static / db-backups)",
        "CI/CD: GitHub Actions · ECR · Argo CD",
        "Auth to AWS: GitHub OIDC",
        "Observability: Prometheus · Grafana · Loki",
        "IaC: Terraform (anime-project-infra)",
    ]:
        add_para(stack.text_frame, "·  " + line, 18, False, NAVY, 10)

    flow = card(s, Inches(6.7), Inches(1.15), Inches(6.1), Inches(5.5), SOFT, TEAL)
    set_text(flow.text_frame, "Data Flow (V2)", 22, True, TEAL)
    for line in [
        "Users → Route53/ACM → ALB Ingress",
        "→ web Pod (Django) → MariaDB Pod",
        "→ EBS PVC (글/데이터 영구 저장)",
        "",
        "사진/미디어 → S3 media/ (DB 아님)",
        "",
        "destroy 후: Git SQL 시드 import",
        "운영 중: CronJob → S3 db-backups/",
        "",
        "eks-stop: PVC 유지",
        "terraform destroy: PVC + S3 삭제",
    ]:
        add_para(flow.text_frame, line, 18, False, NAVY, 8)
    footer(s, 9)


def slide_next(prs):
    s = blank(prs)
    bg(s)
    title_center(s, "향후 계획  ·  3UP")
    cols = [
        (
            ORANGE,
            "Unique Up",
            "차별성",
            [
                "관측 고도화 (Alert · Tempo/OTel)",
                "OIDC 권한 축소 (보안)",
                "백업 덤프 주기·복구 드릴 (DB 안정성)",
                "AIOps (self-healing, auto-remediation) 구현",
            ],
        ),
        (
            GREEN,
            "Complete Up",
            "완성도",
            [
                "Loki/Alloy EKS 반영 완료",
                "시드/목록 검증 CI 게이트",
                "운영 런북·이관 체크리스트",
            ],
        ),
        (
            BLUE,
            "Performance Up",
            "시연·운영",
            [
                "start/stop 데모 시나리오",
                "V1/V2 대비 설명 정리",
                "장애 재현(Missing/SSA) 데모",
            ],
        ),
    ]
    for i, (color, title, sub, lines) in enumerate(cols):
        x = Inches(0.5) + i * Inches(4.25)
        head = card(s, x, Inches(1.15), Inches(4.05), Inches(1.0), color)
        set_text(head.text_frame, title, 22, True, WHITE, PP_ALIGN.CENTER)
        add_para(head.text_frame, sub, 16, False, WHITE, 2, PP_ALIGN.CENTER)
        body = card(s, x, Inches(2.35), Inches(4.05), Inches(4.0), LIGHT)
        body.text_frame.clear()
        pad(body.text_frame)
        first = True
        for line in lines:
            if first:
                set_text(body.text_frame, "·  " + line, 17, False, NAVY)
                first = False
            else:
                add_para(body.text_frame, "·  " + line, 17, False, NAVY, 12)
    footer(s, 10)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    slide_cover(prs)
    slide_v1_def(prs)
    slide_arch_v1(prs)
    slide_v1_to_v2(prs)
    slide_v2_def(prs)
    slide_arch_v2(prs)
    slide_v2_strengths(prs)
    slide_verify(prs)
    slide_stack(prs)
    slide_next(prs)

    out = next_path()
    prs.save(out)
    shutil.copy(out, WORKING)
    print(f"Wrote {out} and {WORKING} ({TOTAL} slides)")


if __name__ == "__main__":
    main()
