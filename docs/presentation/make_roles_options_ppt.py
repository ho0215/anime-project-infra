#!/usr/bin/env python3
"""팀 역할분담 슬라이드 디자인 후보 모음.

출력: ppt/Aniverse_역할분담_옵션.pptx
원하는 옵션(A~F)을 고르면 본편 PPT 19장에 반영하면 됨.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT = "맑은 고딕"
BASE = Path(__file__).resolve().parent
OUT = BASE / "ppt"
OUT.mkdir(parents=True, exist_ok=True)
DST = OUT / "Aniverse_역할분담_옵션.pptx"

SW, SH = 13.333, 7.5
NAVY = RGBColor(15, 23, 42)
BLACK = RGBColor(17, 24, 39)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(107, 114, 128)
LIGHT = RGBColor(243, 244, 246)
SOFT = RGBColor(249, 250, 251)
BLUE = RGBColor(37, 99, 235)
CYAN = RGBColor(14, 165, 233)
ORANGE = RGBColor(249, 115, 22)
PURPLE = RGBColor(139, 92, 246)
TEAL = RGBColor(20, 184, 166)
SLATE = RGBColor(71, 85, 105)

ROLES = [
    (ORANGE, "김현우", "DevOps / GitOps", ["Actions · ECR", "Argo CD", "OIDC · 이관 복구"]),
    (CYAN, "박서이", "EKS · 네트워크 · 보안", ["EKS · VPC", "Ingress · WAF · RBAC", "노드 · NAT"]),
    (TEAL, "김윤주", "컨테이너 · DB · 관측", ["Docker · Helm", "DB Pod · Backup", "Tempo · Alert"]),
    (PURPLE, "강유민", "창작마당 · Compute", ["창작마당", "노드 운영", "ALB · 트래픽"]),
]


def font(run, size=18, bold=False, color=BLACK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = rPr.makeelement(qn(f"a:{tag}"), {})
            rPr.append(el)
        el.set("typeface", FONT)


def pad(tf, l=0.14, t=0.1, r=0.14, b=0.1):
    tf.word_wrap = True
    tf.margin_left = Inches(l)
    tf.margin_right = Inches(r)
    tf.margin_top = Inches(t)
    tf.margin_bottom = Inches(b)


def set_text(tf, text, size=18, bold=False, color=BLACK, align=None):
    tf.clear()
    pad(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    if align is not None:
        p.alignment = align


def add_para(tf, text, size=16, bold=False, color=BLACK, space_before=8, align=None):
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    p.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color=WHITE):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()


def card(slide, x, y, w, h, fill=SOFT, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    pad(sh.text_frame)
    return sh


def title_bar(slide, code, name):
    t = slide.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.55))
    set_text(t.text_frame, f"옵션 {code}  ·  {name}", 28, True, BLACK, PP_ALIGN.CENTER)
    hint = slide.shapes.add_textbox(Inches(0.5), Inches(0.7), Inches(12.3), Inches(0.3))
    set_text(hint.text_frame, "텍스트 카드라 PowerPoint에서 바로 수정 가능합니다", 13, False, GRAY, PP_ALIGN.CENTER)


def footer(slide, code):
    tx = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(12.3), Inches(0.28))
    p = tx.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"Aniverse  ·  역할분담 옵션 {code}"
    font(r, 12, False, GRAY)


# ---------------------------------------------------------------------------
def option_a(prs):
    """컬러 헤더 4열 — 현재 본편과 동일 계열."""
    s = blank(prs)
    bg(s)
    title_bar(s, "A", "컬러 헤더 4열")
    for i, (color, name, area, lines) in enumerate(ROLES):
        x = Inches(0.35) + i * Inches(3.25)
        head = card(s, x, Inches(1.15), Inches(3.05), Inches(1.35), color)
        set_text(head.text_frame, name, 24, True, WHITE, PP_ALIGN.CENTER)
        add_para(head.text_frame, area, 13, False, WHITE, 6, PP_ALIGN.CENTER)
        body = card(s, x, Inches(2.65), Inches(3.05), Inches(4.0), LIGHT)
        body.text_frame.clear()
        pad(body.text_frame)
        first = True
        for line in lines:
            if first:
                set_text(body.text_frame, "·  " + line, 16, False, NAVY, PP_ALIGN.CENTER)
                first = False
            else:
                add_para(body.text_frame, "·  " + line, 16, False, NAVY, 18, PP_ALIGN.CENTER)
    footer(s, "A")


def option_b(prs):
    """왼쪽 액센트 바 + 가로 4행."""
    s = blank(prs)
    bg(s)
    title_bar(s, "B", "가로 행 + 액센트 바")
    for i, (color, name, area, lines) in enumerate(ROLES):
        y = Inches(1.15) + i * Inches(1.4)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), y, Inches(0.12), Inches(1.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()
        c = card(s, Inches(0.75), y, Inches(11.9), Inches(1.2), LIGHT)
        set_text(c.text_frame, f"{name}    ·    {area}", 20, True, BLACK)
        add_para(c.text_frame, "  ·  ".join(lines), 15, False, SLATE, 8)
    footer(s, "B")


def option_c(prs):
    """2×2 그리드."""
    s = blank(prs)
    bg(s)
    title_bar(s, "C", "2×2 그리드")
    positions = [(0.55, 1.2), (6.9, 1.2), (0.55, 4.15), (6.9, 4.15)]
    for (x, y), (color, name, area, lines) in zip(positions, ROLES):
        c = card(s, Inches(x), Inches(y), Inches(5.9), Inches(2.7), SOFT, color)
        set_text(c.text_frame, name, 26, True, color)
        add_para(c.text_frame, area, 15, True, NAVY, 8)
        add_para(c.text_frame, "  ·  ".join(lines), 16, False, SLATE, 12)
    footer(s, "C")


def option_d(prs):
    """번호 뱃지 + 미니멀 리스트."""
    s = blank(prs)
    bg(s)
    title_bar(s, "D", "번호 뱃지 · 미니멀")
    for i, (color, name, area, lines) in enumerate(ROLES):
        y = Inches(1.2) + i * Inches(1.35)
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), y + Inches(0.25), Inches(0.7), Inches(0.7))
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        set_text(badge.text_frame, str(i + 1), 22, True, WHITE, PP_ALIGN.CENTER)
        badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        name_box = s.shapes.add_textbox(Inches(1.7), y, Inches(3.2), Inches(1.15))
        set_text(name_box.text_frame, name, 24, True, BLACK)
        add_para(name_box.text_frame, area, 14, False, color, 6)
        detail = s.shapes.add_textbox(Inches(5.1), y + Inches(0.15), Inches(7.5), Inches(1.0))
        set_text(detail.text_frame, "  ·  ".join(lines), 17, False, NAVY)
    footer(s, "D")


def option_e(prs):
    """영역 칩(pill) + 세로 카드."""
    s = blank(prs)
    bg(s)
    title_bar(s, "E", "영역 칩 + 세로 카드")
    for i, (color, name, area, lines) in enumerate(ROLES):
        x = Inches(0.4) + i * Inches(3.2)
        c = card(s, x, Inches(1.2), Inches(3.0), Inches(5.4), WHITE, color)
        set_text(c.text_frame, name, 26, True, NAVY, PP_ALIGN.CENTER)
        chip = card(s, x + Inches(0.25), Inches(2.35), Inches(2.5), Inches(0.45), color)
        set_text(chip.text_frame, area, 11, True, WHITE, PP_ALIGN.CENTER)
        body = s.shapes.add_textbox(x + Inches(0.2), Inches(3.1), Inches(2.6), Inches(3.2))
        body.text_frame.clear()
        pad(body.text_frame)
        first = True
        for line in lines:
            if first:
                set_text(body.text_frame, "·  " + line, 16, False, NAVY, PP_ALIGN.CENTER)
                first = False
            else:
                add_para(body.text_frame, "·  " + line, 16, False, NAVY, 16, PP_ALIGN.CENTER)
    footer(s, "E")


def option_f(prs):
    """표 형식 깔끔 행."""
    s = blank(prs)
    bg(s)
    title_bar(s, "F", "표 형식")
    # header
    h = card(s, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.6), NAVY)
    set_text(h.text_frame, "이름                    영역                              담당", 16, True, WHITE)
    for i, (color, name, area, lines) in enumerate(ROLES):
        y = Inches(1.95) + i * Inches(1.15)
        fill = LIGHT if i % 2 == 0 else SOFT
        row = card(s, Inches(0.5), y, Inches(12.3), Inches(1.05), fill)
        set_text(row.text_frame, name, 18, True, color)
        add_para(row.text_frame, f"{area}     ·     {'  ·  '.join(lines)}", 14, False, NAVY, 6)
    footer(s, "F")


def option_cover(prs):
    s = blank(prs)
    bg(s, RGBColor(239, 246, 255))
    t = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(2.5))
    set_text(t.text_frame, "팀 역할분담  ·  디자인 옵션", 40, True, BLUE, PP_ALIGN.CENTER)
    add_para(t.text_frame, "A ~ F 중 원하는 스타일을 골라 주세요", 20, False, NAVY, 16, PP_ALIGN.CENTER)
    add_para(
        t.text_frame,
        "모두 PowerPoint에서 글자 수정 가능합니다",
        16,
        False,
        GRAY,
        10,
        PP_ALIGN.CENTER,
    )
    box = s.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(10.3), Inches(1.5))
    set_text(box.text_frame, "A 컬러헤더  ·  B 가로행  ·  C 2×2  ·  D 번호  ·  E 칩카드  ·  F 표", 15, False, SLATE, PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    option_cover(prs)
    option_a(prs)
    option_b(prs)
    option_c(prs)
    option_d(prs)
    option_e(prs)
    option_f(prs)
    prs.save(DST)
    print(f"Wrote {DST} (7 slides: cover + A~F)")


if __name__ == "__main__":
    main()
