#!/usr/bin/env python3
"""하이브리드 발표 PPT — navy/cyan 템플릿 · 큰 글씨 · 구성도 중앙 배치.

구조도: images/hybrid/*.png (make_hybrid_diagrams.py)
출력: ppt/Aniverse_하이브리드_EKS_vN.pptx
"""
from __future__ import annotations

import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FONT_NAME = "맑은 고딕"
BASE = Path(__file__).resolve().parent
IMG = BASE / "images" / "hybrid"
OUT = BASE / "ppt"
OUT.mkdir(parents=True, exist_ok=True)

# Template palette (reference: navy + cyan)
NAVY = RGBColor(10, 17, 40)          # #0A1128
CYAN = RGBColor(0, 174, 239)         # #00AEEF
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(71, 85, 105)
BODY = RGBColor(30, 41, 59)
LIGHT = RGBColor(248, 250, 252)
SOFT = RGBColor(236, 248, 255)
CARD_LINE = RGBColor(186, 230, 253)
GREEN = RGBColor(22, 163, 74)
ORANGE = RGBColor(234, 88, 12)
MUTED = RGBColor(100, 116, 139)

TOTAL = 17
STEM = "Aniverse_하이브리드_EKS"
SW, SH = 13.333, 7.5


def next_version_path() -> Path:
    best = 0
    for p in OUT.glob(f"{STEM}_v*.pptx"):
        m = re.search(r"_v(\d+)\.pptx$", p.name)
        if m:
            best = max(best, int(m.group(1)))
    return OUT / f"{STEM}_v{best + 1}.pptx"


def font(run, size=18, bold=False, color=BODY):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT_NAME
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = rPr.makeelement(qn(f"a:{tag}"), {})
            rPr.append(el)
        el.set("typeface", FONT_NAME)


def pad_tf(tf, left=0.14, top=0.1, right=0.14, bottom=0.1):
    tf.word_wrap = True
    tf.margin_left = Inches(left)
    tf.margin_right = Inches(right)
    tf.margin_top = Inches(top)
    tf.margin_bottom = Inches(bottom)


def set_text(tf, text, size=18, bold=False, color=BODY, align=None):
    tf.clear()
    pad_tf(tf)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    if align is not None:
        p.alignment = align


def add_para(tf, text, size=16, bold=False, color=BODY, space_before=8):
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = text
    font(r, size, bold, color)
    p.space_before = Pt(space_before)
    return p


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.ROUNDED_RECTANGLE
    sh = slide.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.5)
    else:
        sh.line.fill.background()
    pad_tf(sh.text_frame)
    return sh


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def page_bg(slide, fill=WHITE):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    bg.fill.solid()
    bg.fill.fore_color.rgb = fill
    bg.line.fill.background()


def header(slide, title, subtitle=None):
    """White page header: large title + cyan accent line (reference template)."""
    page_bg(slide, WHITE)
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.55))
    set_text(box.text_frame, title, 32, True, NAVY)

    # cyan accent underline
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.88), Inches(1.35), Inches(0.06)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.38))
        set_text(sub.text_frame, subtitle, 16, False, MUTED)


def footer(slide, n):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.15), Inches(SW), Inches(0.35)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = LIGHT
    bar.line.fill.background()
    tx = slide.shapes.add_textbox(Inches(0.55), Inches(7.18), Inches(12.2), Inches(0.28))
    p = tx.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"Aniverse  ·  프로젝트 + EKS 전환"
    font(r, 11, False, MUTED)
    p2 = tx.text_frame.add_paragraph()
    # right page number via separate box
    num = slide.shapes.add_textbox(Inches(11.2), Inches(7.18), Inches(1.6), Inches(0.28))
    p = num.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = f"{n} / {TOTAL}"
    font(r, 11, True, CYAN)
    p.alignment = PP_ALIGN.RIGHT


def put_img(slide, name, top=Inches(1.45), bottom=Inches(7.05), side=0.4):
    """Fit diagram under header — full usable area, centered."""
    path = IMG / name
    if not path.exists():
        return False
    iw, ih = Image.open(path).size
    aspect = ih / float(iw)
    max_w = SW - side * 2
    max_h = bottom.inches - top.inches
    w_in = max_w
    h_in = w_in * aspect
    if h_in > max_h:
        h_in = max_h
        w_in = h_in / aspect
    x_in = (SW - w_in) / 2.0
    y_in = top.inches + (max_h - h_in) / 2.0
    slide.shapes.add_picture(
        str(path), Inches(x_in), Inches(y_in), width=Inches(w_in), height=Inches(h_in)
    )
    return True



# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------
def cover(prs):
    s = blank(prs)
    page_bg(s, NAVY)
    # cyan accent bar top
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CYAN
    bar.line.fill.background()

    t = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(3.2))
    tf = t.text_frame
    set_text(tf, "ANIVERSE", 48, True, WHITE, PP_ALIGN.CENTER)
    add_para(tf, "프로젝트 개요  ·  EKS  ·  GitOps 전환", 22, False, CYAN, 18)
    tf.paragraphs[1].alignment = PP_ALIGN.CENTER
    add_para(tf, "통합 서브컬처 커뮤니티 사이트", 18, False, RGBColor(186, 230, 253), 14)
    tf.paragraphs[2].alignment = PP_ALIGN.CENTER
    add_para(tf, "https://aniverse.my   ·   발표 5~7분", 14, False, MUTED, 22)
    tf.paragraphs[3].alignment = PP_ALIGN.CENTER


def agenda(prs):
    s = blank(prs)
    header(s, "목차", "프로젝트 개요 → 아키텍처 · EKS 전환 → 관측 · 트러블 → 교훈")
    items = [
        ("01", "프로젝트 개요 · 팀 역할"),
        ("02", "아키텍처 v1 · v2 · 왜 EKS"),
        ("03", "랩 → EKS 경로 · 워크로드"),
        ("04", "GitOps · 데이터 백업/복구"),
        ("05", "관측 (Prom / Grafana / Loki)"),
        ("06", "전환 이슈 · 계정 Block / 이관"),
        ("07", "Before / After · 다음"),
    ]
    for i, (num, title) in enumerate(items):
        y = Inches(1.5) + i * Inches(0.74)
        card = rect(s, Inches(0.55), y, Inches(12.2), Inches(0.66), LIGHT, CARD_LINE, True)
        set_text(card.text_frame, f"{num}    {title}", 20, True, NAVY)
        card.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    footer(s, 2)


def service(prs):
    s = blank(prs)
    header(s, "1. 프로젝트 개요", "Aniverse — 통합 서브컬처 사이트")
    goal = rect(s, Inches(0.55), Inches(1.55), Inches(12.2), Inches(1.35), SOFT, CYAN, True)
    set_text(goal.text_frame, "목표", 18, True, CYAN)
    add_para(
        goal.text_frame,
        "애니·굿즈·창작·커뮤니티를 한곳에서 쓰는 통합 서브컬처 사이트를 만드는 것",
        17,
        False,
        BODY,
        10,
    )
    cards = [
        ("제공하는 것", ["장터(굿즈 거래)", "창작 마당", "커뮤니티 · 회원", "https://aniverse.my"]),
        ("왜 만들었나", ["흩어진 서브컬처를 한 플랫폼으로", "웹 + 클라우드 인프라 실습", "배포·운영까지 한 사이클"]),
        ("이 발표에서", ["아키텍처 v1 → v2", "EKS · GitOps 전환", "데이터·관측 · 트러블슈팅"]),
    ]
    for i, (title, lines) in enumerate(cards):
        x = Inches(0.55) + i * Inches(4.15)
        card = rect(s, x, Inches(3.15), Inches(3.95), Inches(3.55), WHITE, CARD_LINE, True)
        set_text(card.text_frame, title, 18, True, CYAN)
        for line in lines:
            add_para(card.text_frame, "·  " + line, 15, False, BODY, 10)
    footer(s, 3)


def roles(prs):
    s = blank(prs)
    header(s, "1. 팀 역할", "윤주=컨테이너·DB·관측  /  현우=Actions·Argo  /  네트워크·컴퓨트")
    put_img(s, "hybrid_00_roles.png")
    footer(s, 4)


def arch_v1(prs):
    s = blank(prs)
    header(s, "2. 아키텍처 v1", "User → Route53/WAF → ALB → ASG/EC2 → RDS · EFS · S3")
    put_img(s, "hybrid_01_arch_v1_vpc.png")
    footer(s, 5)


def arch_v2(prs):
    s = blank(prs)
    header(s, "2. 아키텍처 v2 — EKS", "Ingress/ALB · EKS Pod · DB StatefulSet · ECR · Argo · S3")
    put_img(s, "hybrid_02_arch_v2_vpc.png")
    footer(s, 6)


def why_eks(prs):
    s = blank(prs)
    header(s, "2. 왜 EKS · DB Pod 인가", "멘토링 · 비용 · 동작 확인 우선")
    rows = [
        ("배포 · 재현", "이미지 + GitOps로 동일 환경 재현"),
        ("스케일", "노드 스케일 + 파드 HPA를 분리해 설계"),
        ("DB Pod", "RDS 상시 비용 대신 StatefulSet + PVC (학습/2차)"),
        ("원칙", "아키텍처 확정 → 동작 확인 → 관측·보안 고도화"),
    ]
    for i, (t, b) in enumerate(rows):
        y = Inches(1.55) + i * Inches(1.25)
        card = rect(s, Inches(0.55), y, Inches(12.2), Inches(1.1), SOFT, CYAN, True)
        set_text(card.text_frame, t, 20, True, CYAN)
        add_para(card.text_frame, b, 17, False, BODY, 8)
    footer(s, 7)


def migration(prs):
    s = blank(prs)
    header(s, "3. 랩 → EKS 전환 경로", "Compose → Lab K8s(Helm) → EKS + Argo")
    put_img(s, "hybrid_03_migration_path.png")
    footer(s, 8)


def workloads(prs):
    s = blank(prs)
    header(s, "3. 앱 · DB 워크로드", "윤주: Docker · Helm · StatefulSet — EKS web/db 1/1")
    put_img(s, "hybrid_04_workloads.png")
    footer(s, 9)


def gitops(prs):
    s = blank(prs)
    header(s, "4. GitOps · CI/CD", "Actions → ECR(sha-*) → values bump → Argo sync · OIDC")
    put_img(s, "hybrid_05_gitops.png")
    footer(s, 10)


def data(prs):
    s = blank(prs)
    header(s, "4. 데이터 — 백업 · 복구", "restore Job · CronJob → S3  ·  media는 S3 sync")
    put_img(s, "hybrid_06_data.png")
    footer(s, 11)


def observability(prs):
    s = blank(prs)
    header(s, "5. 관측", "완료: Prom+Grafana · Loki/Alloy  —  예정: AlertManager · Tempo/OTel")
    put_img(s, "hybrid_07_observability.png")
    footer(s, 12)


def ops(prs):
    s = blank(prs)
    header(s, "5. 운영 포인트", "HTTPS · DNS · 비용 start/stop")
    items = [
        ("HTTPS", "ACM ISSUED 후 Ingress ALB · Route53"),
        ("비용", "EKS start/stop — 미사용 시 노드·NAT 절감"),
        ("보안", "장기 Access Key 금지 · GitHub OIDC · Pod S3 IRSA"),
        ("검증", "health 200 + 시드/목록 UI (Actions 초록만으로 끝내지 않음)"),
    ]
    for i, (t, b) in enumerate(items):
        col, row = i % 2, i // 2
        x = Inches(0.55) + col * Inches(6.25)
        y = Inches(1.55) + row * Inches(2.55)
        card = rect(s, x, y, Inches(6.0), Inches(2.3), WHITE, CARD_LINE, True)
        set_text(card.text_frame, t, 22, True, CYAN)
        add_para(card.text_frame, b, 17, False, BODY, 16)
    footer(s, 13)


def issue_tech(prs):
    s = blank(prs)
    header(s, "6. 전환 이슈 — 기술", "같은 Missing이라도 원인이 둘")
    left = rect(s, Inches(0.55), Inches(1.55), Inches(5.95), Inches(5.2), WHITE, ORANGE, True)
    set_text(left.text_frame, "A. Job sync-wave", 22, True, ORANGE)
    for line in [
        "Ingress가 ALB ADDRESS 전 Progressing",
        "다음 wave Job이 영구 Missing",
        "조치: Job에 높은 wave 금지",
        "DB ready는 Job 내부 ping",
    ]:
        add_para(left.text_frame, "·  " + line, 16, False, BODY, 14)
    right = rect(s, Inches(6.8), Inches(1.55), Inches(5.95), Inches(5.2), WHITE, CYAN, True)
    set_text(right.text_frame, "B. ServerSideApply", 22, True, CYAN)
    for line in [
        "terminatingReplicas ComparisonError",
        "live SSA 잔존 시 재발",
        "조치: SSA 제거 · client-side apply",
        "install 후 syncOptions strip",
    ]:
        add_para(right.text_frame, "·  " + line, 16, False, BODY, 14)
    footer(s, 14)


def issue_account(prs):
    s = blank(prs)
    header(s, "6. 계정 Block → 이관", "표면은 start 실패 · 근본은 키 유출")
    put_img(s, "hybrid_08_account.png")
    footer(s, 15)


def before_after(prs):
    s = blank(prs)
    header(s, "7. Before / After", "v1 EC2 운영 → v2 EKS GitOps")
    left = rect(s, Inches(0.55), Inches(1.55), Inches(5.95), Inches(5.2), LIGHT, MUTED, True)
    set_text(left.text_frame, "Before (v1)", 24, True, MUTED)
    for line in [
        "EC2 + CodeDeploy",
        "RDS MariaDB",
        "수동·에이전트 배포",
        "환경 재현 어려움",
    ]:
        add_para(left.text_frame, "·  " + line, 18, False, BODY, 18)
    right = rect(s, Inches(6.8), Inches(1.55), Inches(5.95), Inches(5.2), SOFT, CYAN, True)
    set_text(right.text_frame, "After (v2)", 24, True, CYAN)
    for line in [
        "EKS + Argo CD",
        "DB Pod + PVC · SQL/CronJob 백업",
        "Actions → ECR → GitOps",
        "Prom/Grafana · Loki 관측",
    ]:
        add_para(right.text_frame, "·  " + line, 18, False, BODY, 18)
    footer(s, 16)


def closing(prs):
    s = blank(prs)
    header(s, "7. 배운 점 · 다음", "하이브리드 발표 마무리")
    lessons = [
        (
            "배운 점",
            [
                "아키텍처 확정 후 동작 확인이 1순위",
                "장기 키 금지 → OIDC · 이관 시 ARN 전수 교체",
                "Actions 초록 ≠ 시드/목록 검증",
            ],
        ),
        (
            "다음",
            [
                "Loki/Alloy EKS 반영 · AlertManager",
                "Tempo + Django OpenTelemetry",
                "OIDC 권한 축소 · Budgets",
            ],
        ),
    ]
    for i, (title, lines) in enumerate(lessons):
        x = Inches(0.55) + i * Inches(6.25)
        card = rect(s, x, Inches(1.55), Inches(6.0), Inches(5.2), WHITE, CARD_LINE, True)
        set_text(card.text_frame, title, 24, True, CYAN)
        for line in lines:
            add_para(card.text_frame, "·  " + line, 17, False, BODY, 18)
    footer(s, 17)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    cover(prs)
    agenda(prs)
    service(prs)
    roles(prs)
    arch_v1(prs)
    arch_v2(prs)
    why_eks(prs)
    migration(prs)
    workloads(prs)
    gitops(prs)
    data(prs)
    observability(prs)
    ops(prs)
    issue_tech(prs)
    issue_account(prs)
    before_after(prs)
    closing(prs)

    out = next_version_path()
    prs.save(out)
    print(f"Wrote {out} ({TOTAL} slides)")


if __name__ == "__main__":
    main()
